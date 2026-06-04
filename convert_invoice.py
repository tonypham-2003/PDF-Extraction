"""
convert_invoice.py - PDF Invoice to Excel
Usage:
    python3 convert_invoice.py              # skip files already converted
    python3 convert_invoice.py --force      # overwrite all
    python3 convert_invoice.py --dry-run    # preview only
"""
import re, sys, argparse, time, pdfplumber, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

BASE_DIR   = Path(__file__).parent
INPUT_DIR  = BASE_DIR / "Input"
OUTPUT_DIR = BASE_DIR / "Output"

COL_WIDTHS = [6, 14, 14, 45, 10, 12, 14, 12, 14, 20, 18]

# Regex compiled once at module level
RE_VNDEC      = re.compile(r'(VNDEC\d+)')
RE_INV_NO     = re.compile(r'No\.?[:\s]+([A-Z0-9\-]+(?:-TP)?)', re.I)
RE_COMMODITY  = re.compile(r'Commodity Code[^\n]*\n(\d+)')
RE_CMA_BLOCK  = re.compile(
    r'Size/Type\s+Charge Description.*?\n(.*?)(?:Rate of Exchange|VAT applied)',
    re.DOTALL | re.I
)
RE_CMA_FULL   = re.compile(
    r'^(\S+)\s+C\s+(.+?)\s+(\d+)([A-Z]+)\s+([\d,]+\.\d+)([A-Z]+)\s+([\d,]+\.\d+)\s+([\d,]+\.\d+)$'
)
RE_CMA_SHORT  = re.compile(
    r'^(\S+)\s+C\s+(.+?)\s+([\d,]+\.\d+)([A-Z]+)\s+([\d,]+\.\d+)\s+([\d,]+\.\d+)$'
)
RE_CMA_ROW    = re.compile(r'^(40|20|45)')
RE_TAX_SUFFIX = re.compile(r'\s+[A-Z][0-9]\s*$')
RE_COMMERCIAL = re.compile(
    r'^(\d+)\s+(.+?)\s+([A-Za-z]+)\s+(\d[\d,]*)\s+([\d,]+)\s+([\d,]+)$'
)

# Border cached
_BORDER = None
def _thin():
    global _BORDER
    if _BORDER is None:
        s = Side(style="thin")
        _BORDER = Border(left=s, right=s, top=s, bottom=s)
    return _BORDER


def extract_text(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def parse_pdf(pdf_path):
    text = extract_text(pdf_path)
    if not text.strip():
        return pdf_path.stem, "", [], "LOW"

    inv_match = RE_VNDEC.search(text) or RE_INV_NO.search(text)
    invoice_no = inv_match.group(1) if inv_match else None
    comm_match = RE_COMMODITY.search(text)
    commodity_code = comm_match.group(1) if comm_match else ""

    items = []
    block_match = RE_CMA_BLOCK.search(text)
    if block_match:
        items = _parse_cma(block_match.group(1).strip(), commodity_code)
    if not items:
        items = _parse_commercial(text)

    return invoice_no or pdf_path.stem, commodity_code, items, _confidence(invoice_no, items)


def _parse_cma(block, commodity_code):
    items = []
    stt = 1
    for line in block.splitlines():
        line = line.strip()
        if not line or not RE_CMA_ROW.match(line):
            continue
        m = RE_CMA_FULL.match(line)
        if m:
            _, desc, bq, bu, rate, curr, _a, vnd = m.groups()
            items.append(_row(stt, RE_TAX_SUFFIX.sub('', desc).strip(),
                               commodity_code, bq, bu, rate + " " + curr, vnd))
            stt += 1; continue
        m2 = RE_CMA_SHORT.match(line)
        if m2:
            _, desc, rate, curr, _a, vnd = m2.groups()
            items.append(_row(stt, RE_TAX_SUFFIX.sub('', desc).strip(),
                               commodity_code, "", "", rate + " " + curr, vnd))
            stt += 1
    return items


def _parse_commercial(text):
    items = []
    for line in text.splitlines():
        m = RE_COMMERCIAL.match(line.strip())
        if m:
            stt, desc, uom, qty, price, subtotal = m.groups()
            items.append(_row(int(stt), desc.strip(), "", qty, uom, price, subtotal))
    return items


def _row(stt, ten_hang, ma, so_luong_1, don_vi_tinh_1, don_gia, tong_tri_gia):
    return {"stt": stt, "ma_hang": ma, "ma_hs": ma, "ten_hang": ten_hang,
            "xuat_xu": "", "so_luong_1": so_luong_1, "don_vi_tinh_1": don_vi_tinh_1,
            "so_luong_2": "", "don_vi_tinh_2": "", "don_gia": don_gia, "tong_tri_gia": tong_tri_gia}


def _confidence(invoice_no, items):
    if not items or not invoice_no:
        return "LOW"
    for item in items:
        if not item.get("ten_hang") or not item.get("tong_tri_gia"):
            return "LOW"
    return "HIGH"


def write_excel(items, out_path, invoice_no):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sach hang"
    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = "DANH SACH HANG - " + invoice_no
    c.font = Font(bold=True, size=13, color="1F4E79")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    hfill = PatternFill("solid", fgColor="1F4E79")
    hfont = Font(bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    right  = Alignment(horizontal="right",  vertical="center")

    col_labels = ["STT","Ma hang","Ma HS","Ten hang","Xuat xu",
                  "So luong 1","Don vi tinh 1","So luong 2","Don vi tinh 2",
                  "Don gia","Tong tri gia"]
    for ci, h in enumerate(col_labels, 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = hfont; cell.fill = hfill
        cell.alignment = center; cell.border = _thin()
    ws.row_dimensions[2].height = 30

    for ri, item in enumerate(items, 3):
        vals = [item["stt"], item["ma_hang"], item["ma_hs"], item["ten_hang"],
                item["xuat_xu"], item["so_luong_1"], item["don_vi_tinh_1"],
                item["so_luong_2"], item["don_vi_tinh_2"], item["don_gia"], item["tong_tri_gia"]]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font = Font(size=10); cell.border = _thin()
            cell.alignment = left if ci == 4 else (right if ci in (10, 11) else center)
        ws.row_dimensions[ri].height = 35

    for ci, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A3"
    wb.save(out_path)


def out_name(inv): return "danh_sach_hang_" + inv + ".xlsx"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--force",   action="store_true")
    args = parser.parse_args()

    pdfs = sorted(INPUT_DIR.glob("*.pdf"))
    if not pdfs:
        print("Khong co file PDF nao trong Input.")
        sys.exit(0)

    OUTPUT_DIR.mkdir(exist_ok=True)

    if not args.dry_run and not args.force:
        new_pdfs = [p for p in pdfs if not (OUTPUT_DIR / out_name(
            (RE_VNDEC.search(p.stem) or re.search(r'(.+)', p.stem)).group(1)
        )).exists()]
        skipped = len(pdfs) - len(new_pdfs)
        if skipped:
            print("Bo qua " + str(skipped) + " file da co output (them --force de overwrite)")
        pdfs = new_pdfs

    if not pdfs:
        print("Tat ca da duoc convert.")
        return

    print("Xu ly " + str(len(pdfs)) + " file:")
    for f in pdfs: print("  - " + f.name)

    t0 = time.perf_counter()
    results = [(p,) + parse_pdf(p) for p in pdfs]
    t_parse = time.perf_counter()

    if args.dry_run:
        print("\n[DRY-RUN - chua tao file]")
        for pdf_path, invoice_no, _, items, conf in results:
            print("\n" + pdf_path.name + "  Invoice: " + (invoice_no or "?") + "  [" + conf + "]")
            if items:
                for it in items:
                    print("  " + str(it['stt']) + ". " + str(it['ten_hang']) +
                          "  |  " + str(it['don_gia']) + "  |  " + str(it['tong_tri_gia']))
            else:
                print("  (Khong doc duoc du lieu)")
        has_low = any(c == "LOW" for _,_,_,_,c in results)
        print("\nCANH BAO: Confidence thap!" if has_low else "\nTat ca confidence CAO.")
        return

    converted = 0
    for pdf_path, invoice_no, _, items, conf in results:
        print("\n" + pdf_path.name + "  [" + conf + "]")
        if not items:
            print("  Khong doc duoc du lieu - bo qua.")
            continue
        if conf == "LOW":
            for it in items:
                print("  " + str(it['stt']) + ". " + str(it['ten_hang']) + " | " + str(it['don_gia']) + " | " + str(it['tong_tri_gia']))
            ans = input("  Confidence THAP - tao Excel khong? (y/N): ").strip().lower()
            if ans != "y":
                continue
        write_excel(items, OUTPUT_DIR / out_name(invoice_no), invoice_no)
        print("  OK: " + out_name(invoice_no) + " (" + str(len(items)) + " dong)")
        converted += 1

    t_end = time.perf_counter()
    print("\nHoan tat! " + str(converted) + " file"
          + "  |  Parse: " + str(round((t_parse - t0)*1000)) + "ms"
          + "  |  Total: " + str(round((t_end - t0)*1000)) + "ms")


if __name__ == "__main__":
    main()
