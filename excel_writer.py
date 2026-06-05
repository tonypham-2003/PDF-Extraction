"""excel_writer.py — write extracted items to .xlsx following template.xls layout.

Template structure (matches template.xls exactly):
  Row 1 : column headers  (blue background, white bold text)
  Row 2+: data rows
  Columns: STT | Mã hàng | Mã HS | Tên hàng | Xuất xứ |
           Số lượng 1 | Đơn vị tính 1 | Số lượng 2 | Đơn vị tính 2 |
           Đơn giá | Tổng trị giá

Number format : US input "2,349.00" → Vietnamese output "2.349,00" (all numeric strings)
Đơn giá       → stored as string in Vietnamese format ("2.349,00")
Tổng trị giá  → stored as float with #,##0.00 format (renders per Windows locale)
Uncertain     → AI returns "uncertain":[field,...] per row; those cells are highlighted yellow for manual review
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_COLOR   = "1F4E79"
UNCERTAIN_COLOR = "FFFF00"   # yellow — needs manual review
COL_WIDTHS     = [6, 20, 14, 60, 10, 13, 16, 13, 16, 18, 18]
HEADERS = [
    "STT", "Mã hàng ", "Mã HS", "Tên hàng", "Xuất xứ",
    "Số lượng 1", "Đơn vị tính 1", "Số lượng 2", "Đơn vị tính 2",
    "Đơn giá", "Tổng trị giá",
]

_BORDER = None
def _thin():
    global _BORDER
    if _BORDER is None:
        s = Side(style="thin")
        _BORDER = Border(left=s, right=s, top=s, bottom=s)
    return _BORDER


def _us_to_vn(s: str) -> str:
    """Convert number string from US format (2,349.00) to Vietnamese (2.349,00).
    Handles plain numbers too: "2349" → "2.349,00", "25.5" → "25,50".
    """
    if not s:
        return s
    s = str(s).strip()
    try:
        val = float(s.replace(",", ""))
        us_fmt = f"{val:,.2f}"                                   # "2,349.00"
        vn_fmt = us_fmt.replace(",", "X").replace(".", ",").replace("X", ".")  # "2.349,00"
        return vn_fmt
    except ValueError:
        return s


def _to_number(val):
    """Convert a value to float for Tổng trị giá (numeric cell)."""
    if val is None or val == "":
        return ""
    s = str(val).lstrip("?").strip()
    try:
        return float(s.replace(",", ""))
    except (ValueError, TypeError):
        return val




def write_excel(items: list[dict], out_path) -> None:
    """Write items to xlsx following template.xls structure."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # ── Row 1: headers ────────────────────────────────────────────────────────
    hfill      = PatternFill("solid", fgColor=HEADER_COLOR)
    hfont      = Font(bold=True, color="FFFFFF", size=10)
    ufill      = PatternFill("solid", fgColor=UNCERTAIN_COLOR)
    center     = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border     = _thin()

    for ci, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hfont; cell.fill = hfill
        cell.alignment = center; cell.border = border
    ws.row_dimensions[1].height = 30

    # ── Data rows from row 2 ──────────────────────────────────────────────────
    left  = Alignment(horizontal="left",  vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    for stt, item in enumerate(items, 1):
        ri = stt + 1

        don_gia_raw   = item.get("don_gia",      "")
        tong_raw      = item.get("tong_tri_gia", "")

        # Convert numeric strings to Vietnamese format
        don_gia_vn = _us_to_vn(don_gia_raw)
        tong_float = _to_number(tong_raw)       # stays float for numeric cell

        vals = [
            stt,
            item.get("ma_hang",       ""),
            item.get("ma_hs",         ""),
            item.get("ten_hang",      ""),
            item.get("xuat_xu",       ""),
            _us_to_vn(item.get("so_luong_1",    "")),
            item.get("don_vi_tinh_1", ""),
            _us_to_vn(item.get("so_luong_2",    "")),
            item.get("don_vi_tinh_2", ""),
            don_gia_vn,
            tong_float,
        ]

        # Build set of column indices to highlight yellow (from AI's "uncertain" list)
        _field_to_col = {
            "ma_hang": 2, "ma_hs": 3, "ten_hang": 4, "xuat_xu": 5,
            "so_luong_1": 6, "don_vi_tinh_1": 7,
            "so_luong_2": 8, "don_vi_tinh_2": 9,
            "don_gia": 10, "tong_tri_gia": 11,
        }
        uncertain_cols = {
            _field_to_col[f]
            for f in item.get("uncertain", [])
            if f in _field_to_col
        }

        for ci, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font   = Font(size=10)
            cell.border = border
            cell.alignment = (left  if ci == 4 else
                              right if ci in (10, 11) else center)
            if ci == 11 and isinstance(val, float):
                cell.number_format = "#,##0.00"
            if ci in uncertain_cols:
                cell.fill = ufill

        ws.row_dimensions[ri].height = 35

    for ci, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = "A2"
    wb.save(out_path)
